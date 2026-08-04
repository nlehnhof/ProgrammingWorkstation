"""Header-name-based access to the per-airport gate spreadsheets.

Every airport sheet is the same shape (see `devices/*/IP_TEMPLATE.xlsx`):

    PBB Gate | Gate IP | Netmask | Gateway | PBB SN | Jetway PN | MAC Address |
    Router Programmed On | Label | Crash Report | PRG # | Router Tested On |
    Crash Report | Test # | Notes

The device scripts used to reach into these by *position* -- `row[idx - 3]` for
the IP, `column=7` for the MAC. That silently returns the wrong cell the moment
someone inserts a column, and it did: the same lookup produced different fields
in different sheets. Everything here resolves a column by the text in its
header row instead, so re-ordering a sheet cannot change what gets read or
written.

Note "Crash Report" legitimately appears twice (programming crash, then test
crash), which is why `get_header_map` returns a *list* of columns per name and
`find_column` takes an `occurrence`.
"""

import os
from contextlib import contextmanager

import openpyxl

from resources.utilities.network_utils import is_valid_ip, validate_subnet

# Logical field -> header text in the sheet. Device code refers to the logical
# name so a header rename is a one-line change here.
GATE_IP = "Gate IP"
NETMASK = "Netmask"
GATEWAY = "Gateway"

# Cells that mean "no value" despite not being empty: rows written by earlier
# tooling carry the literal text "None" in the PN/MAC columns, and a label must
# never print the word "None" onto a physical sticker.
PLACEHOLDER_CELLS = {"", "none", "nan", "n/a", "na", "null", "-", "--"}


@contextmanager
def open_workbook(path):
    """Yield a workbook and guarantee it closes, even if the body raises.

    openpyxl holds the file open until `close()`. Without this the sheet stays
    locked after any error, and the next run fails to save with a
    PermissionError that looks like the operator left Excel open.

    A file that is not really a spreadsheet becomes a ValueError naming it --
    `devices/TR/JKC-SLC.xlsx` is a real example, and openpyxl's raw
    `BadZipFile: File is not a zip file` gives an operator nothing to act on.
    """
    try:
        workbook = openpyxl.load_workbook(path)
    except Exception as exc:
        if isinstance(exc, OSError):
            raise
        raise ValueError(f"{path} is not a readable .xlsx file: {exc}") from exc

    try:
        yield workbook
    finally:
        workbook.close()


def get_header_map(sheet, header_row=1):
    """Map header text to the 1-based column numbers carrying that text.

    Duplicated headers accumulate, so {"Crash Report": [10, 13]}.
    Raises ValueError if the header row is empty -- better than handing back
    an empty map that turns every later lookup into a confusing KeyError.
    """
    headers = {}
    for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row), ()):
        if cell.value is None:
            continue
        name = str(cell.value).strip()
        if name:
            headers.setdefault(name, []).append(cell.column)

    if not headers:
        raise ValueError(f"Row {header_row} of the sheet has no column headers.")
    return headers


def find_column(header_map, name, occurrence=0):
    """1-based column number for `name`, matched case-insensitively.

    `occurrence` picks between repeated headers (0 = the first "Crash Report",
    1 = the second). Raises KeyError naming what was actually available, which
    is the difference between a fixable error message and a bare KeyError.
    """
    columns = header_map.get(name)
    if columns is None:
        wanted = name.strip().lower()
        for key, value in header_map.items():
            if key.strip().lower() == wanted:
                columns = value
                break

    if not columns:
        raise KeyError(f"No column named {name!r}. Found: {sorted(header_map)}")
    if occurrence >= len(columns):
        raise KeyError(
            f"Sheet has {len(columns)} column(s) named {name!r}; "
            f"occurrence {occurrence} was requested."
        )
    return columns[occurrence]


def require_columns(header_map, names):
    """Fail once, listing every missing header, instead of one at a time."""
    present = {key.strip().lower() for key in header_map}
    missing = [name for name in names if name.strip().lower() not in present]
    if missing:
        raise ValueError(
            f"Sheet is missing required column(s): {', '.join(missing)}. "
            f"Found: {sorted(header_map)}"
        )


def find_row_by_value(sheet, value, min_row=2):
    """Row number of the first cell equal to `value`, or None.

    Compared as text, because the gate identifier arrives from a GUI dropdown
    as a string while the sheet stores it as a number.
    """
    wanted = str(value)
    for row in sheet.iter_rows(min_row=min_row):
        for cell in row:
            if cell.value is not None and str(cell.value) == wanted:
                return cell.row
    return None


def lookup_excel(path, gate):
    """Network settings for one gate: {"gate_ip", "netmask", "gateway"}.

    Raises rather than returning blanks, so a bad sheet stops the run before
    any hardware is touched:
      * ValueError -- required headers absent, or the address/netmask pair in
        the sheet is not a usable one.
      * KeyError   -- that gate is not in this sheet.
    """
    with open_workbook(path) as workbook:
        sheet = workbook.active
        if sheet is None:
            raise ValueError(f"{path} has no active worksheet.")

        header_map = get_header_map(sheet)
        require_columns(header_map, [GATE_IP, NETMASK, GATEWAY])

        row = find_row_by_value(sheet, gate)
        if row is None:
            raise KeyError(f"Gate {gate!r} was not found in {os.path.basename(path)}.")

        def value_at(name):
            cell = sheet.cell(row=row, column=find_column(header_map, name))
            return None if cell.value is None else str(cell.value).strip()

        gate_ip = value_at(GATE_IP)
        netmask = value_at(NETMASK)
        gateway = value_at(GATEWAY)

    if not is_valid_ip(gate_ip):
        raise ValueError(f"Gate {gate} has an unusable IP address in the sheet: {gate_ip!r}")

    ok, network = validate_subnet(gate_ip, netmask)
    if not ok:
        raise ValueError(
            f"Gate {gate}: {gate_ip} does not sit inside the subnet its netmask "
            f"{netmask!r} describes ({network})."
        )

    return {"gate_ip": gate_ip, "netmask": netmask, "gateway": gateway}


def cell_text(value):
    """Sheet cell -> text safe to print on a label ("" for placeholders)."""
    text = "" if value is None else str(value).strip()
    return "" if text.lower() in PLACEHOLDER_CELLS else text


# ---------------------------------------------------------------------------
# GUI dropdown helpers
# ---------------------------------------------------------------------------

def get_excel_files(folder_path):
    """Airport sheets in a device folder, for the Airport dropdown."""
    if not os.path.isdir(folder_path):
        return []
    return sorted(
        name for name in os.listdir(folder_path)
        if name.endswith(".xlsx")
        and not name.startswith("~$")       # Excel's lock files
        and name != "oshkosh_log.xlsx"
    )


def get_dropdown(filepath, column_name="PBB SN"):
    """Gate identifiers from one airport sheet, for the Gate dropdown.

    Reads the serial column by header name so an inserted column cannot shift
    the dropdown onto some unrelated field.
    """
    try:
        with open_workbook(filepath) as workbook:
            sheet = workbook.active
            if sheet is None:
                return []
            column = find_column(get_header_map(sheet), column_name)
            return [
                str(row[0].value)
                for row in sheet.iter_rows(min_row=2, min_col=column, max_col=column)
                if row[0].value is not None
            ]
    except (OSError, ValueError, KeyError) as exc:
        print(f"Could not read gate options from {filepath}: {exc}", flush=True)
        return []
