from unittest.mock import MagicMock, patch

import openpyxl
import pytest

from resources.utilities.excel_utils import (
    find_column,
    find_row_by_value,
    get_header_map,
    lookup_excel,
    open_workbook,
    require_columns,
)

REAL_HEADERS = [
    "PBB Gate", "Gate IP", "Netmask", "Gateway", "PBB SN", "Jetway PN",
    "MAC Address", "Router Programmed On", "Label", "Crash Report",
    "PRG #", "Router Tested On", "Crash Report", "Test #", "Notes",
]


def make_workbook(path, rows, headers=REAL_HEADERS):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    wb.save(path)
    return path


def test_get_header_map_matches_real_template_layout(tmp_path):
    path = make_workbook(tmp_path / "wb.xlsx", rows=[])
    wb = openpyxl.load_workbook(path)
    header_map = get_header_map(wb.active)
    wb.close()

    assert header_map["Gate IP"] == [2]
    assert header_map["Netmask"] == [3]
    assert header_map["Gateway"] == [4]
    assert header_map["MAC Address"] == [7]
    # "Crash Report" is duplicated (programming crash log + test crash log)
    assert header_map["Crash Report"] == [10, 13]


def test_get_header_map_raises_on_blank_header_row(tmp_path):
    path = tmp_path / "wb.xlsx"
    wb = openpyxl.Workbook()
    wb.save(path)
    loaded = openpyxl.load_workbook(path)
    with pytest.raises(ValueError):
        get_header_map(loaded.active)
    loaded.close()


def test_require_columns_raises_with_missing_names_listed():
    header_map = {"Gate IP": [2], "Netmask": [3]}
    with pytest.raises(ValueError, match="Gateway"):
        require_columns(header_map, ["Gate IP", "Netmask", "Gateway"])


def test_find_column_handles_duplicate_headers_by_occurrence():
    header_map = {"Crash Report": [10, 13]}
    assert find_column(header_map, "Crash Report", occurrence=0) == 10
    assert find_column(header_map, "Crash Report", occurrence=1) == 13


def test_find_column_raises_keyerror_for_missing_name():
    with pytest.raises(KeyError):
        find_column({}, "Gate IP")


def test_find_row_by_value_locates_gate_row(tmp_path):
    path = make_workbook(
        tmp_path / "wb.xlsx",
        rows=[
            [21, "10.28.18.2", "255.255.255.0", "0.0.0.0", 10000],
            [22, "10.120.30.4", "255.255.255.0", "0.0.0.0", 10001],
        ],
    )
    wb = openpyxl.load_workbook(path)
    assert find_row_by_value(wb.active, "10001") == 3
    assert find_row_by_value(wb.active, "no-such-gate") is None
    wb.close()


def test_lookup_excel_returns_gate_ip_netmask_gateway_by_header_not_position(tmp_path):
    path = make_workbook(
        tmp_path / "wb.xlsx",
        rows=[[21, "10.28.18.2", "255.255.255.0", "10.28.18.1", 10000]],
    )
    result = lookup_excel(str(path), "10000")
    assert result == {"gate_ip": "10.28.18.2", "netmask": "255.255.255.0", "gateway": "10.28.18.1"}


def test_lookup_excel_still_correct_when_columns_are_reordered(tmp_path):
    # Column order differs from the "real" template, but headers are still
    # named correctly -- a positional lookup would silently return the
    # wrong values here, which is exactly the bug this fixes.
    reordered_headers = ["PBB SN", "Gateway", "Gate IP", "Netmask", "PBB Gate"]
    path = make_workbook(
        tmp_path / "wb.xlsx",
        rows=[[10000, "10.28.18.1", "10.28.18.2", "255.255.255.0", 21]],
        headers=reordered_headers,
    )
    result = lookup_excel(str(path), "10000")
    assert result == {"gate_ip": "10.28.18.2", "netmask": "255.255.255.0", "gateway": "10.28.18.1"}


def test_lookup_excel_raises_when_gate_missing(tmp_path):
    path = make_workbook(
        tmp_path / "wb.xlsx",
        rows=[[21, "10.28.18.2", "255.255.255.0", "10.28.18.1", 10000]],
    )
    with pytest.raises(KeyError):
        lookup_excel(str(path), "99999")


def test_lookup_excel_raises_when_required_headers_missing(tmp_path):
    path = make_workbook(
        tmp_path / "wb.xlsx",
        rows=[[21, 10000]],
        headers=["PBB Gate", "PBB SN"],
    )
    with pytest.raises(ValueError):
        lookup_excel(str(path), "10000")


def test_lookup_excel_raises_on_invalid_ip(tmp_path):
    path = make_workbook(
        tmp_path / "wb.xlsx",
        rows=[[21, "not-an-ip", "255.255.255.0", "10.28.18.1", 10000]],
    )
    with pytest.raises(ValueError):
        lookup_excel(str(path), "10000")


def test_open_workbook_closes_even_when_body_raises():
    fake_wb = MagicMock()
    with patch("resources.utilities.excel_utils.openpyxl.load_workbook", return_value=fake_wb):
        with pytest.raises(RuntimeError):
            with open_workbook("dummy.xlsx") as wb:
                assert wb is fake_wb
                raise RuntimeError("boom")
    fake_wb.close.assert_called_once()
